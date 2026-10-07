"""
Excel 文件读写模块


职责：

    - 将采集结果写入 Excel 文件
    - 从 Excel 文件读取全部数据
    - 支持 dict / dataclass对象
    - 自动创建输出目录


调用关系：

    FileIoManager
            |
            v
        ExcelIO
            |
            v
        XLSX文件
"""

import logging
from dataclasses import asdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)


class ExcelIO:
    """
    Excel 读写器
    """

    def __init__(
        self,
        *,
        output_path,
        filename="data.xlsx",
        sheet_name="Sheet1",
    ):
        """
        初始化 Excel 写入器


        参数：

            output_path:

                输出目录


            filename:

                Excel文件名


            sheet_name:

                工作表名称
        """

        self.output_path = Path(output_path)

        self.filename = filename

        self.file_path = self.output_path / self.filename

        self.sheet_name = sheet_name

        self._create_output_dir()

    def _create_output_dir(self):
        """
        创建输出目录
        """

        self.output_path.mkdir(
            parents=True,
            exist_ok=True,
        )

    def write(
        self,
        data,
    ):
        """
        写入一条数据


        支持：

            dict

            dataclass对象
        """

        row = self._convert_data(data)

        if not row:
            logger.warning("Excel写入数据为空")

            return

        self._write_row(row)

    def read(
        self,
    ):
        """
        读取全部数据


        返回：

            list[dict]

                每行一个 dict，键为表头


        文件不存在或工作表为空时返回空列表
        """

        if not self.file_path.exists():
            logger.warning(
                "Excel文件不存在: %s",
                self.file_path,
            )

            return []

        workbook = load_workbook(
            self.file_path,
            read_only=True,
        )

        if self.sheet_name not in workbook.sheetnames:
            logger.warning(
                "工作表不存在: %s",
                self.sheet_name,
            )

            return []

        worksheet = workbook[self.sheet_name]

        rows = worksheet.iter_rows(
            values_only=True,
        )

        # 跳过开头整行为空的行，找到真正的表头

        # （兼容旧写入器或外部工具留下的首行空行）

        headers = None

        for first_row in rows:
            if first_row is None:
                continue

            if any(value is not None for value in first_row):
                headers = first_row

                break

        if not headers:
            return []

        result = []

        for row in rows:
            if row is None:
                continue

            # 跳过整行空值

            if all(value is None for value in row):
                continue

            result.append(
                dict(
                    zip(
                        headers,
                        row,
                    )
                )
            )

        return result

    def _convert_data(
        self,
        data,
    ):
        """
        数据转换
        """

        # dataclass对象

        if hasattr(
            data,
            "__dataclass_fields__",
        ):
            return asdict(data)

        # dict

        if isinstance(
            data,
            dict,
        ):
            return data

        raise TypeError(f"Excel不支持的数据类型: {type(data)}")

    def _write_row(
        self,
        row,
    ):
        """
        写入Excel行数据
        """

        if self.file_path.exists():
            workbook = load_workbook(self.file_path)

        else:
            workbook = Workbook()

        if self.sheet_name in workbook.sheetnames:
            worksheet = workbook[self.sheet_name]

        else:
            worksheet = workbook.create_sheet(self.sheet_name)

        # 判断是否需要写表头

        # 新建 Workbook 的工作表为空，直接写表头；

        # 已有文件则按第一行是否已有内容判断

        # （访问 .cell() 会实例化单元格，使 append 落到下一行，

        # 因此仅在加载已有文件时探测单元格）

        if self.file_path.exists():
            has_header = (
                worksheet.max_row >= 1
                and worksheet.cell(
                    1,
                    1,
                ).value
                is not None
            )

        else:
            has_header = False

        # 仅当无表头时写表头

        if not has_header:
            worksheet.append(list(row.keys()))

        # 数据写入

        worksheet.append(list(row.values()))

        workbook.save(self.file_path)

        logger.debug(
            "Excel写入成功: %s",
            self.file_path,
        )
